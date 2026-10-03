#!/usr/bin/env python3
"""Extract the CSG Fantasy Football Sheet's consensus columns -> csg_consensus.json.

CSG is a **complementary** source, never the valuation. `build_tool_data.py` owns
`worth`/`vbd`, recomputed from your scraped ESPN scoring (see CLAUDE.md: "Valuation is
league-accurate — do not regress this"). What CSG adds is the OUTSIDE VIEW: where the
wider market ranks and prices a player, so the console can show where your league-accurate
number disagrees with consensus. Disagreement is the whole point — it is where the edge is.

    CSG*-auction.xlsm  "Overall" tab  ──►  csg_consensus.json  ──►  build_tool_data.py
                                                                     (merged as p["mkt"])

Layout (stable across v11-v14): the "Overall" tab carries a control block in rows 1-10,
the header on row 11, and players from row 12. Columns are grouped into the sections CSG
documents: rankings (ESPN/NFL/Yahoo/ECR), prices/ADP, Boris Chen tiers, value skews,
BeerSheets, ElBoberto VBD, and Gold Score.

COVERAGE IS UNEVEN and varies by year — always check the printed table rather than
assuming a column is populated. Known gaps as of the 2026 v14.1 sheet:
  - BeerSheets (BS Val / PosScarcity / BeerTier) is 0/385 — it needs a manual paste into
    the hidden "Beersheet Paste" tab and BeerSheets had not shipped 2026 values yet.
    It IS populated for 2023-2025, so history works.
  - "AdjVBD" exists only from v14.1 (2026); earlier sheets have no such column.
  - Gold Score is deliberately top-of-board only (~45 players).
  - NFL rankings are unpopulated in 2026.

Run:   python3 draft_sheets/extract_csg.py          (all CSG workbooks found)
       python3 pipeline.py csg                      (same, via the pipeline)
"""
import glob
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "csg_consensus.json")

HEADER_ROW = 11          # the "Overall" tab's player-table header
FIRST_DATA_ROW = 12

# CSG header -> our key. Only what is decision-useful; internal match/helper columns are
# dropped. Left side must match the sheet's row-11 text exactly.
FIELDS = {
    "Position":       "pos",
    "Team":           "team",
    "Bye":            "bye",
    "Age":            "age",
    "Status":         "status",        # injury note, e.g. "Q: Week 1 (Knee)"
    # § rankings
    "ESPN":           "espn_rank",     # what your ESPN-using leaguemates see
    "Yahoo":          "yahoo_rank",
    "ECR":            "ecr",           # FantasyPros expert consensus
    # § prices / ADP
    "ESPN Price":     "espn_price",
    "Average Price":  "mkt_price",     # cross-site average auction $ — the market number
    "Proj Price":     "csg_proj_price",
    # § Boris Chen
    "Boris":          "boris_tier",
    # § value skews
    "ECR Skew":       "ecr_skew",
    # § BeerSheets
    "BS Val":         "bs_val",
    "PosScarcity":    "pos_scarcity",
    "BeerTier":       "beer_tier",
    # § ElBoberto VBD (CSG's own implementation — NOT our league-accurate vbd)
    "VBD":            "csg_vbd",
    "AdjVBD":         "csg_adj_vbd",   # v14.1+ only; scarcity-adjusted
    "VAL%":           "val_pct",
    # § Gold Score
    "Gold Rank":      "gold",          # a score, not a rank: higher = better
}

# values that mean "no data" in this sheet
BLANK = {"", "na", "n/a", "-", "--", "none"}


# The single name-join key lives in fftiers.names (shared with the manage board);
# re-exported here so `from extract_csg import norm_name` keeps working.
try:
    from fftiers.names import norm_name
except ImportError:  # standalone use without the fftiers install
    def norm_name(s):
        s = str(s or "").lower()
        s = s.replace("&", "and")
        s = re.sub(r"[.'`,]", "", s)
        s = re.sub(r"\b(jr|sr|ii|iii|iv|v)\b", "", s)
        s = re.sub(r"[^a-z0-9]+", " ", s)
        return " ".join(s.split())


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    t = str(v).strip()
    if t.lower() in BLANK:
        return None
    try:
        return float(t.replace("$", "").replace(",", ""))
    except ValueError:
        return None


def year_of(path):
    m = re.search(r"-\s*(20\d{2})\s*v", os.path.basename(path)) or \
        re.search(r"(20\d{2})", os.path.basename(path))
    return m.group(1) if m else os.path.basename(path)


def read_overall(path):
    """Parse one CSG workbook's Overall tab into {norm_name: {field: value}}."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Overall" not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"{os.path.basename(path)}: no 'Overall' tab — not a CSG sheet?")
    ws = wb["Overall"]
    rows = list(ws.iter_rows(min_row=HEADER_ROW, max_row=ws.max_row, values_only=True))
    wb.close()
    if not rows:
        return {}, {}
    hdr = rows[0]
    col = {}
    for i, name in enumerate(hdr):
        if name is None:
            continue
        key = FIELDS.get(str(name).strip())
        if key and key not in col:      # first occurrence wins
            col[key] = i
    if "Player" not in [str(h).strip() for h in hdr if h]:
        raise SystemExit(f"{os.path.basename(path)}: header row {HEADER_ROW} has no 'Player'")
    pi = [i for i, h in enumerate(hdr) if h and str(h).strip() == "Player"][0]

    text_fields = {"pos", "team", "status"}
    out = {}
    for r in rows[1:]:
        raw = r[pi] if pi < len(r) else None
        if not raw or not str(raw).strip():
            continue
        rec = {"name": str(raw).strip()}
        for key, i in col.items():
            v = r[i] if i < len(r) else None
            if key in text_fields:
                t = str(v).strip() if v is not None else ""
                if t and t.lower() not in BLANK:
                    rec[key] = t
            else:
                n = _num(v)
                # CSG writes 0 for "not priced"/"unranked" in these columns, not a real 0
                if n is not None and not (n == 0 and key not in ("bye", "ecr_skew")):
                    rec[key] = round(n, 3)
        out[norm_name(raw)] = rec
    return out, col


def main():
    paths = sorted(glob.glob(os.path.join(HERE, "CSG*auction.xlsm")))
    if not paths:
        raise SystemExit(
            "No CSG*auction.xlsm found in draft_sheets/. Drop the CSG sheet(s) there "
            "(they stay local — .gitignore keeps third-party sheets out of git)."
        )
    out = {}
    for p in paths:
        players, col = read_overall(p)
        yr = year_of(p)
        out[yr] = players
        # coverage audit — the point of this print is that gaps are the norm, so a build
        # never silently relies on an empty column.
        keys = [k for k in FIELDS.values() if k not in ("pos", "team", "bye", "age")]
        cov = {k: sum(1 for r in players.values() if r.get(k) is not None) for k in keys}
        print(f"{yr}: {len(players)} players  ({os.path.basename(p)})")
        for k in keys:
            n = cov[k]
            if n == 0:
                flag = "  <- EMPTY in this sheet"
            elif n < len(players) * 0.5:
                flag = "  <- partial"
            else:
                flag = ""
            print(f"      {k:16} {n:>4}/{len(players)}{flag}")
    json.dump(out, open(OUT, "w"), indent=1, sort_keys=True)
    print(f"\nwrote {OUT}  ({len(out)} seasons)")
    print("  build_tool_data.py merges the CURRENT season onto each console player as "
          "p['mkt']; it never touches worth/vbd.")


if __name__ == "__main__":
    sys.exit(main())
